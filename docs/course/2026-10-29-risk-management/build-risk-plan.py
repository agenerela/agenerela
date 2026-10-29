"""Builds the Agenerela risk management plan workbook for the COMP 490 in-class practice."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1] if len(sys.argv) > 1 else "Agenerela-Risk-Management-Plan.xlsx"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
SUB_FILL = PatternFill("solid", start_color="DDEBF7")
INPUT_FILL = PatternFill("solid", start_color="FFFDE7")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


RED, ORANGE, YELLOW, GREEN, LGREEN = "F8696B", "FFB366", "FFEB84", "63BE7B", "C6EFCE"

TEAM = [
    "Yevhen Mishchenko", "Hunter Hudson", "Timothy Brustinov",
    "Tigran Kolsuzyan", "Hero Jaiyen", "Maxim Goloubitsky", "Whole team",
]
RISK_TYPES = ["Estimation", "Organizational", "People", "Requirements", "Technology", "Tools"]
AFFECTS = ["Project", "Product", "Business", "Project + Product"]
PROB = [("Very low", 1, "< 10%"), ("Low", 2, "10–25%"), ("Moderate", 3, "25–50%"),
        ("High", 4, "50–75%"), ("Very high", 5, "> 75%")]
EFFECT = [("Insignificant", 1, "Barely noticed; absorbed inside the sprint"),
          ("Tolerable", 2, "Costs days of work or one demo feature; plan holds"),
          ("Serious", 3, "Misses a phase deadline or weakens a headline claim"),
          ("Catastrophic", 4, "Project cannot be delivered or its core claim fails")]
STRATEGIES = ["Avoidance", "Minimization", "Contingency"]
STATUSES = ["Monitoring", "Observed", "Resolved"]

# (type, risk, identified by, affects, impact areas, probability, effect,
#  what-if, strategy, action plan, indicators, status)
RISKS = [
    ("People",
     "The core (interfaces, pipeline, schema, guards) is built mostly by one person; if the lead is "
     "sick or overloaded at a deadline, every other area is blocked.",
     "Yevhen Mishchenko", "Project", "Phase 3–4 schedule; all research areas", "Moderate", "Serious",
     "What if the lead is out for a week right before the Phase 3 due date (9 Nov)?",
     "Minimization",
     "Every core PR is reviewed by a second member, so at least one other person knows each part. "
     "Design decisions written as decision records in the build plan. Pair one teammate with the lead on "
     "the guard pipeline.",
     "Core PRs merging with no reviewer; issues waiting on the lead for more than a week.",
     "Monitoring"),
    ("Technology",
     "Small local models that fit an 8 GB GPU pick the wrong action too often for the framework to be "
     "convincing in a game.",
     "Yevhen Mishchenko", "Product", "Decision accuracy; demo games; final presentation", "High", "Serious",
     "What if accuracy stays near the prototype's 53–64% even after the guards are in?",
     "Contingency",
     "Keep the provider swappable (Ollama, cloud providers). Guards (confidence check, validation) turn wrong "
     "answers into safe 'no action'. Phase 4 harness measures every change with a control arm; if local "
     "models stay weak, demo with a larger or cloud model and report the local numbers honestly.",
     "Phase 4 A/B runs below target; many 'wrong-but-legal' outcomes in telemetry.",
     "Monitoring"),
    ("Requirements",
     "The way observations are given to the model (templated lines vs summaries) changes after the Agent API "
     "is built, forcing rework of the decision request.",
     "Hunter Hudson", "Project", "Agent / DecideAsync API; observation sources (#91, #122)", "Moderate",
     "Tolerable",
     "What if the Phase 4 A/B shows summaries win and the API only accepts plain strings?",
     "Avoidance",
     "Agree the observation-source interface with the area owner before the core builds it. Keep "
     "observations behind an interface so a new strategy is an implementation, not an API change.",
     "Interface changes requested after merge; open questions on #122.",
     "Monitoring"),
    ("People",
     "Key members are unavailable at critical times: midterms, finals, Thanksgiving and winter break fall "
     "around phase deadlines.",
     "Hunter Hudson", "Project", "Phase 3 (9 Nov) and Phase 4 deadlines; reviews", "High", "Tolerable",
     "What if two reviewers are out during finals week and PRs pile up?",
     "Minimization",
     "Each member posts known unavailable dates on the board at the start of each sprint. Every research "
     "area has a backup reviewer. Plan lighter sprints for exam weeks.",
     "Review queue growing; members missing stand-ups; issues idle more than a week.",
     "Monitoring"),
    ("Estimation",
     "The effort to write and label 200+ evaluation prompts (with a second annotator) is underestimated, "
     "delaying Phase 4.",
     "Timothy Brustinov", "Project + Product", "Evaluation set (#85); Phase 4 harness; every accuracy claim",
     "High", "Serious",
     "What if only ~60 prompts are labelled when Phase 4 is due?",
     "Minimization",
     "Split labelling across all six members with a weekly quota, using the Prompt Set window / labelling "
     "artifact. Track count per week. Measure on the prompts that exist and grow n rather than waiting.",
     "Labelled count flat for a sprint; annotator labels still pending (PR #118).",
     "Observed"),
    ("Tools",
     "Other model providers (Gemini, OpenAI, in-process) support different parts of JSON Schema, so a schema "
     "that works on Ollama fails elsewhere.",
     "Timothy Brustinov", "Product", "Provider layer; schema serializer", "Moderate", "Tolerable",
     "What if a provider rejects enums or required fields the framework depends on (e.g. Gemini rejects "
     "empty enum strings)?",
     "Avoidance",
     "Keep the schema to the common subset (no empty enums, required 'target' with 'no_target'). Add a "
     "provider-specific test per provider before claiming support.",
     "Provider tests failing; schema features added that only one provider accepts.",
     "Monitoring"),
    ("Technology",
     "An Ollama upgrade changes how structured output or token probabilities behave, silently changing "
     "results or breaking the confidence guard.",
     "Tigran Kolsuzyan", "Product", "Schema constraint; confidence guard; benchmark numbers", "Low",
     "Serious",
     "What if accuracy shifts between two runs only because Ollama was updated in between?",
     "Avoidance",
     "Pin and record the Ollama version in every measurement. Re-run schema_visibility_probe.py after "
     "each upgrade. Never compare runs across versions.",
     "Probe output differs from findings.md; unexplained jumps in benchmark results.",
     "Monitoring"),
    ("Estimation",
     "The size of the software is underestimated: framework, editor tooling and three demo games are more "
     "than two semesters allow.",
     "Tigran Kolsuzyan", "Project", "Demo games (#47); editor tooling (Phase 5); final review", "Moderate",
     "Serious",
     "What if by March only one demo game runs end to end?",
     "Contingency",
     "Planned cut order from the build plan: cut art polish first, then drop GreyBox2D. Never cut "
     "GreyBoxStrategy (proves the genre-agnostic claim) or Phase 4.",
     "Phase dates slipping by more than a sprint; demo milestones missed after each phase.",
     "Monitoring"),
    ("Technology",
     "Agent memory grows the prompt until it exceeds the model's context or makes decisions too slow for "
     "gameplay on an 8 GB card.",
     "Hero Jaiyen", "Product", "Memory strategies (#119); latency; VRAM", "Moderate", "Serious",
     "What if a long play session makes each decision take several seconds?",
     "Minimization",
     "Give memory a size budget and a summarising/forgetting strategy. Log prompt length and latency per "
     "decision in telemetry; measure one model resident at a time.",
     "Prompt tokens or latency rising over a session; model partly off the GPU (ollama ps).",
     "Monitoring"),
    ("Technology",
     "The model chooses a legal but wrong target (right action, wrong character or object), which the "
     "validator cannot catch.",
     "Hero Jaiyen", "Product", "Target grounding (#120); player trust in demos", "Moderate", "Serious",
     "What if the companion attacks the wrong enemy in front of the reviewers?",
     "Minimization",
     "Only offer reachable / visible targets in the schema. Track 'wrong-but-legal' as its own outcome "
     "class. Add evaluation prompts with several similar targets.",
     "Wrong-but-legal share rising in evaluation; target bug reports from demo play-tests.",
     "Monitoring"),
    ("Tools",
     "Unity and Git do not work together smoothly across six people: missing or regenerated .meta files, "
     "merge conflicts in scenes and assets.",
     "Maxim Goloubitsky", "Project", "Every PR; demo game scenes", "Moderate", "Tolerable",
     "What if a merge breaks references in a demo scene the day before a presentation?",
     "Avoidance",
     "Hygiene CI and pre-commit hook (python tools/ci/hygiene.py). One owner per scene. Restore committed "
     ".meta files, never commit regenerated GUIDs.",
     "Red 'Repo hygiene' checks; 'missing script' warnings in Unity.",
     "Monitoring"),
    ("Estimation",
     "The rate of defect repair is underestimated: PRs merged without review let bugs and tests that cannot "
     "fail reach the test branch.",
     "Maxim Goloubitsky", "Project + Product", "Code quality; time spent on follow-up fixes", "High",
     "Tolerable",
     "What if a hidden bug in the pipeline invalidates a week of measurements?",
     "Minimization",
     "Review before merge (REVIEWING.md); prove each new test can fail once. Keep a follow-up issue for "
     "every defect found after merge.",
     "Follow-up fix PRs after merges; review findings on already-merged PRs.",
     "Observed"),
    ("Organizational",
     "Course requirements or deadlines change (new deliverables, presentations), taking time away from "
     "planned phases.",
     "Whole team", "Business", "Grades; schedule", "Moderate", "Tolerable",
     "What if a large course deliverable lands in the same week as a phase gate?",
     "Minimization",
     "Check Canvas weekly; split course deliverables across members so the phase owner keeps building.",
     "New Canvas assignments not in the plan; phase work idle during course weeks.",
     "Monitoring"),
    ("Tools",
     "Not every member has a GPU strong enough to run models locally, so they cannot run live tests or "
     "measurements.",
     "Whole team", "Project", "RequiresOllama tests; benchmarks; asset server", "Low", "Serious",
     "What if the one capable machine fails in the week before a demo?",
     "Contingency",
     "EditMode tests run without a model. Use small models or a cloud provider for development; keep "
     "measurements on one documented machine; keep a second machine set up as backup.",
     "Members unable to run PlayMode tests; asset server unreachable.",
     "Monitoring"),
]

wb = Workbook()

# ---------------- Legend sheet (built first so the register can reference it) ----------------
lg = wb.active
lg.title = "Legend"
reg = wb.create_sheet("Risk Register", 0)

lg["A1"] = "Scales and how to use this workbook"
lg["A1"].font = f(True, size=14)
lg["A3"] = "Probability"
lg["E3"] = "Effects (seriousness)"
for c in ("A3", "E3"):
    lg[c].font = f(True, size=11)
for col, h in zip("ABC", ["Level", "Score", "Likelihood"]):
    lg[f"{col}4"] = h
for col, h in zip("EFG", ["Level", "Score", "Meaning"]):
    lg[f"{col}4"] = h
for row in (lg["A4:C4"][0] + lg["E4:G4"][0]):
    row.font = f(True, "FFFFFF")
    row.fill = HEAD_FILL
    row.border = BORDER
for i, (lvl, sc, txt) in enumerate(PROB):
    r = 5 + i
    lg[f"A{r}"], lg[f"B{r}"], lg[f"C{r}"] = lvl, sc, txt
for i, (lvl, sc, txt) in enumerate(EFFECT):
    r = 5 + i
    lg[f"E{r}"], lg[f"F{r}"], lg[f"G{r}"] = lvl, sc, txt
for rng in ("A5:C9", "E5:G8"):
    for row in lg[rng]:
        for c in row:
            c.font = f()
            c.border = BORDER
            c.alignment = WRAP

lg["A11"] = "Priority = Probability score × Effect score (1–20)"
lg["A11"].font = f(True, size=11)
for col, h in zip("ABC", ["Priority", "Min score", "Action"]):
    lg[f"{col}12"] = h
    lg[f"{col}12"].font = f(True, "FFFFFF")
    lg[f"{col}12"].fill = HEAD_FILL
    lg[f"{col}12"].border = BORDER
prio = [("High", 12, "Act now; owner reports on it every sprint", RED),
        ("Medium", 6, "Plan in place; check every sprint", YELLOW),
        ("Low", 1, "Watch the indicators", GREEN)]
for i, (p, mn, act, col) in enumerate(prio):
    r = 13 + i
    lg[f"A{r}"], lg[f"B{r}"], lg[f"C{r}"] = p, mn, act
    for c in ("A", "B", "C"):
        lg[f"{c}{r}"].font = f()
        lg[f"{c}{r}"].border = BORDER
        lg[f"{c}{r}"].alignment = WRAP
    lg[f"A{r}"].fill = fill(col)
lg["D13"] = "Thresholds are the team's choice; change B13:B14 and every priority updates."
lg["D13"].font = f(italic=True, color="595959")

lg["A17"] = "Strategies"
lg["A17"].font = f(True, size=11)
strat_txt = [("Avoidance", "Reduce the probability that the risk arises."),
             ("Minimization", "Reduce the impact if it does arise."),
             ("Contingency", "Be prepared for the worst; have a plan ready to switch to.")]
for i, (s, t) in enumerate(strat_txt):
    lg[f"A{18+i}"], lg[f"B{18+i}"] = s, t
    lg[f"A{18+i}"].font = f(True)
    lg[f"B{18+i}"].font = f()

lg["A22"] = "Status"
lg["A22"].font = f(True, size=11)
stat_txt = [("Monitoring", "Not happened; indicators are watched.", "FFF2CC"),
            ("Observed", "Indicators seen or risk partly happened; plan is being carried out.", "F4B183"),
            ("Resolved", "Risk gone or plan finished; kept for the record.", LGREEN)]
for i, (s, t, col) in enumerate(stat_txt):
    lg[f"A{23+i}"], lg[f"B{23+i}"] = s, t
    lg[f"A{23+i}"].font = f(True)
    lg[f"A{23+i}"].fill = fill(col)
    lg[f"B{23+i}"].font = f()

lg["A27"] = "How to edit"
lg["A27"].font = f(True, size=11)
howto = [
    "Edit only the pale-yellow cells on 'Risk Register'. Score and Priority are formulas; do not type over them.",
    "Pick Type, Affects, Probability, Effects, Strategy and Status from the drop-downs; colours update by "
    "conditional formatting.",
    "Add a new risk on the next empty row (up to row 60); copy the Score and Priority formulas down if needed.",
    "Once per sprint: review every row, change probability/effects/status, set 'Last reviewed', and add a line "
    "to 'Sprint Review Log'.",
]
for i, t in enumerate(howto):
    lg[f"A{28+i}"] = f"{i+1}. {t}"
    lg[f"A{28+i}"].font = f()
lg.column_dimensions["A"].width = 16
lg.column_dimensions["B"].width = 10
lg.column_dimensions["C"].width = 44
lg.column_dimensions["D"].width = 3
lg.column_dimensions["E"].width = 16
lg.column_dimensions["F"].width = 8
lg.column_dimensions["G"].width = 52
# Strategy/status text sits in B and must not be clipped by C
for r in list(range(18, 21)) + list(range(23, 26)):
    lg.merge_cells(f"B{r}:G{r}")
for r in range(28, 32):
    lg.merge_cells(f"A{r}:G{r}")
lg.merge_cells("D13:G14")
lg["D13"].alignment = WRAP

# ---------------- Risk Register ----------------
reg["A1"] = "Agenerela — Risk Management Plan"
reg["A1"].font = f(True, size=16, color="1F4E78")
reg["A2"] = ("Team project: Agenerela (Unity framework for LLM-driven game agents), CSUN COMP 490 · "
             "Created 29 Oct 2026 · Reviewed once per sprint")
reg["A2"].font = f(italic=True, color="595959")

headers = ["ID", "Risk type", "Risk", "Identified by / owner", "Affects", "Impact area(s)",
           "Probability", "Effects", "Score", "Priority", "What-if question", "Strategy",
           "Planned actions", "Indicators to watch", "Status", "Last reviewed"]
widths = [6, 14, 46, 18, 12, 26, 12, 13, 7, 10, 34, 14, 56, 34, 12, 12]
HR = 4
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = reg.cell(row=HR, column=i, value=h)
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    reg.column_dimensions[c.column_letter].width = w
reg.row_dimensions[HR].height = 30

FIRST = HR + 1
LAST_FMT = 60  # formatting/validation reach, leaving room for new risks
for i, r in enumerate(RISKS):
    row = FIRST + i
    (rtype, risk, who, aff, area, p, e, whatif, strat, plan, ind, status) = r
    vals = [f"R{i+1:02d}", rtype, risk, who, aff, area, p, e, None, None, whatif, strat, plan, ind, status,
            "2026-10-29"]
    for col, v in enumerate(vals, start=1):
        if v is not None:
            reg.cell(row=row, column=col, value=v)

for row in range(FIRST, LAST_FMT + 1):
    reg[f"I{row}"] = (f'=IF(OR(G{row}="",H{row}=""),"",'
                      f'INDEX(Legend!$B$5:$B$9,MATCH(G{row},Legend!$A$5:$A$9,0))*'
                      f'INDEX(Legend!$F$5:$F$8,MATCH(H{row},Legend!$E$5:$E$8,0)))')
    reg[f"J{row}"] = (f'=IF(I{row}="","",IF(I{row}>=Legend!$B$13,"High",'
                      f'IF(I{row}>=Legend!$B$14,"Medium","Low")))')
    for col in range(1, len(headers) + 1):
        c = reg.cell(row=row, column=col)
        c.font = f()
        c.border = BORDER
        c.alignment = WRAP
        if col not in (9, 10):
            c.fill = INPUT_FILL
    for col in (1, 5, 7, 8, 9, 10, 12, 15, 16):
        reg.cell(row=row, column=col).alignment = CENTER
    reg[f"I{row}"].font = f(True)
    reg[f"J{row}"].font = f(True)
for i in range(len(RISKS)):
    reg.row_dimensions[FIRST + i].height = 84

reg.freeze_panes = "D5"
reg.auto_filter.ref = f"A{HR}:P{LAST_FMT}"
reg["I4"].comment = Comment("Probability score × Effect score, from the Legend sheet.", "Team")
reg["J4"].comment = Comment("High ≥ Legend!B13, Medium ≥ Legend!B14, else Low.", "Team")

# Drop-downs
def dv(options, rng):
    v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    v.add(rng)
    reg.add_data_validation(v)

dv(RISK_TYPES, f"B{FIRST}:B{LAST_FMT}")
dv(TEAM, f"D{FIRST}:D{LAST_FMT}")
dv(AFFECTS, f"E{FIRST}:E{LAST_FMT}")
dv([p[0] for p in PROB], f"G{FIRST}:G{LAST_FMT}")
dv([e[0] for e in EFFECT], f"H{FIRST}:H{LAST_FMT}")
dv(STRATEGIES, f"L{FIRST}:L{LAST_FMT}")
dv(STATUSES, f"O{FIRST}:O{LAST_FMT}")

# Conditional formatting
def cf_eq(rng, value, color, bold=False):
    reg.conditional_formatting.add(
        rng, CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(color),
                        font=Font(name=FONT, bold=bold)))

R = lambda col: f"{col}{FIRST}:{col}{LAST_FMT}"
for v, col in [("Very high", RED), ("High", ORANGE), ("Moderate", YELLOW), ("Low", LGREEN),
               ("Very low", GREEN)]:
    cf_eq(R("G"), v, col)
for v, col in [("Catastrophic", "C00000"), ("Serious", RED), ("Tolerable", YELLOW),
               ("Insignificant", GREEN)]:
    cf_eq(R("H"), v, col)
for v, col in [("High", RED), ("Medium", YELLOW), ("Low", GREEN)]:
    cf_eq(R("J"), v, col, bold=True)
for v, col in [("Monitoring", "FFF2CC"), ("Observed", "F4B183"), ("Resolved", LGREEN)]:
    cf_eq(R("O"), v, col)
# Score cell shaded with its priority
reg.conditional_formatting.add(R("I"), FormulaRule(formula=[f'$J{FIRST}="High"'], fill=fill(RED)))
reg.conditional_formatting.add(R("I"), FormulaRule(formula=[f'$J{FIRST}="Medium"'], fill=fill(YELLOW)))
reg.conditional_formatting.add(R("I"), FormulaRule(formula=[f'$J{FIRST}="Low"'], fill=fill(GREEN)))
# Resolved rows fade to grey text
reg.conditional_formatting.add(
    f"A{FIRST}:N{LAST_FMT}",
    FormulaRule(formula=[f'$O{FIRST}="Resolved"'], font=Font(name=FONT, color="A6A6A6", strike=True)))

reg.sheet_view.zoomScale = 90
reg.page_setup.orientation = "landscape"
reg.page_setup.fitToWidth = 1
reg.page_setup.fitToHeight = 0
reg.sheet_properties.pageSetUpPr.fitToPage = True
reg.print_title_rows = f"{HR}:{HR}"

# ---------------- Risk Matrix ----------------
mx = wb.create_sheet("Risk Matrix", 1)
mx["A1"] = "Risk matrix — number of open risks (Monitoring or Observed) in each cell"
mx["A1"].font = f(True, size=14)
mx["A2"] = "Counts are formulas over 'Risk Register' and update when probability, effects or status change."
mx["A2"].font = f(italic=True, color="595959")
mx["B4"] = "Effects →"
mx["B4"].font = f(True)
mx["A5"] = "Probability ↓"
mx["A5"].font = f(True, "FFFFFF")
mx["A5"].fill = HEAD_FILL
mx["A5"].alignment = CENTER
eff_cols = "BCDE"
for j, (lvl, sc, _) in enumerate(EFFECT):
    c = mx[f"{eff_cols[j]}5"]
    c.value = lvl
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
reg_rng = lambda col: f"'Risk Register'!${col}${FIRST}:${col}${LAST_FMT}"
for i, (plvl, psc, _) in enumerate(reversed(PROB)):
    r = 6 + i
    mx[f"A{r}"] = plvl
    mx[f"A{r}"].font = f(True)
    mx[f"A{r}"].fill = SUB_FILL
    mx[f"A{r}"].border = BORDER
    mx[f"A{r}"].alignment = CENTER
    for j, (elvl, esc, _) in enumerate(EFFECT):
        c = mx[f"{eff_cols[j]}{r}"]
        c.value = (f'=COUNTIFS({reg_rng("G")},$A{r},{reg_rng("H")},{eff_cols[j]}$5,{reg_rng("O")},"Monitoring")'
                   f'+COUNTIFS({reg_rng("G")},$A{r},{reg_rng("H")},{eff_cols[j]}$5,{reg_rng("O")},"Observed")')
        score = psc * esc
        c.fill = fill(RED if score >= 12 else YELLOW if score >= 6 else GREEN)
        c.font = f(True, size=14)
        c.alignment = CENTER
        c.border = BORDER
    mx.row_dimensions[r].height = 36
mx.column_dimensions["A"].width = 16
for col in eff_cols:
    mx.column_dimensions[col].width = 16
mx.column_dimensions["D"].width = 22

mx["A12"] = "Summary"
mx["A12"].font = f(True, size=11)
summary = [
    ("Total risks", f'=COUNTA({reg_rng("A")})'),
    ("High priority", f'=COUNTIF({reg_rng("J")},"High")'),
    ("Medium priority", f'=COUNTIF({reg_rng("J")},"Medium")'),
    ("Low priority", f'=COUNTIF({reg_rng("J")},"Low")'),
    ("Observed", f'=COUNTIF({reg_rng("O")},"Observed")'),
    ("Resolved", f'=COUNTIF({reg_rng("O")},"Resolved")'),
]
for i, (lbl, fml) in enumerate(summary):
    r = 13 + i
    mx[f"A{r}"], mx[f"B{r}"] = lbl, fml
    mx[f"A{r}"].font = f()
    mx[f"B{r}"].font = f(True)
    mx[f"A{r}"].border = mx[f"B{r}"].border = BORDER

mx["D12"] = "Risks per member"
mx["D12"].font = f(True, size=11)
for i, name in enumerate(TEAM):
    r = 13 + i
    mx[f"D{r}"] = name
    mx[f"E{r}"] = f'=COUNTIF({reg_rng("D")},D{r})'
    mx[f"D{r}"].font = f()
    mx[f"E{r}"].font = f(True)
    mx[f"D{r}"].border = mx[f"E{r}"].border = BORDER
mx.conditional_formatting.add(
    "E13:E18", CellIsRule(operator="lessThan", formula=["2"], fill=fill(RED)))
mx["D20"] = "Red = fewer than two risks."
mx["D20"].font = f(italic=True, color="595959")

# ---------------- Sprint Review Log ----------------
lgv = wb.create_sheet("Sprint Review Log")
lgv["A1"] = "Sprint review log — one line per review"
lgv["A1"].font = f(True, size=14)
cols = [("Date", 12), ("Sprint", 10), ("Reviewed by", 30), ("New risks", 18),
        ("Changed (probability / effects / status)", 44), ("Ground-rule changes", 44)]
for i, (h, w) in enumerate(cols, start=1):
    c = lgv.cell(row=3, column=i, value=h)
    c.font = f(True, "FFFFFF")
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
    lgv.column_dimensions[c.column_letter].width = w
first_entry = ["2026-10-29", "In class", "All six members", "R01–R14",
               "Initial register. R05 and R12 already Observed.",
               "Proposed: review before merge; post unavailable dates each sprint; weekly labelling quota."]
for i, v in enumerate(first_entry, start=1):
    c = lgv.cell(row=4, column=i, value=v)
    c.font = f()
    c.alignment = WRAP
    c.border = BORDER
for r in range(5, 25):
    for col in range(1, 7):
        c = lgv.cell(row=r, column=col)
        c.border = BORDER
        c.fill = INPUT_FILL
        c.font = f()
        c.alignment = WRAP
lgv.row_dimensions[4].height = 45

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT)
